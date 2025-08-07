import io
from docx import Document
from PyPDF2 import PdfReader
from langchain_core.documents import Document as LangchainDocument

from utils.logger import logger


def extract_text(data, filename):
    document_text = ""
    if filename.endswith('.txt'):
        document_text = data.decode('utf-8')
    elif filename.endswith('.pdf'):
        document_text = extract_text_from_pdf(data)
    elif filename.endswith('.docx'):
        document_text = extract_text_from_docx(data)
    return LangchainDocument(page_content=document_text, metadata={"filename": filename})



def extract_text_from_pdf(data):
    with io.BytesIO(data) as pdf_file:
        reader = PdfReader(pdf_file)
        text = ""
        for page in reader.pages:
            text += page.extract_text()
    return text

def extract_text_from_docx(data):
    try:
        with io.BytesIO(data) as docx_file:
            doc = Document(docx_file)
            text = '\n'.join([para.text for para in doc.paragraphs])
        return text
    except Exception as e:
        logger.error(f"An error occurred: {e}")
        return None
